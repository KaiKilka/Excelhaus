import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

DUNKELBLAU  = "1F3864"
WEISS       = "FFFFFF"

def titelzeile(ws, titel, spalten_headers, tab_farbe, col_widths=None):
    ws.sheet_properties.tabColor = tab_farbe
    n = len(spalten_headers)
    ws.merge_cells(f"A1:{get_column_letter(n)}1")
    c = ws["A1"]
    c.value = titel
    c.font = Font(bold=True, color=WEISS, size=12)
    c.fill = PatternFill("solid", fgColor=DUNKELBLAU)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    for i, h in enumerate(spalten_headers, 1):
        cell = ws.cell(row=2, column=i, value=h)
        cell.font = Font(bold=True, color=WEISS, size=10)
        cell.fill = PatternFill("solid", fgColor=DUNKELBLAU)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 22
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A2:{get_column_letter(n)}2"
    ws.freeze_panes = "A3"

wb = openpyxl.Workbook()

# ── 1_PSP ─────────────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "1_PSP"
titelzeile(ws1,
    "PROJEKTSTRUKTURPLAN – Kernsanierung EFH Baujahr 1969",
    ["PSP-Code","Ebene","Gewerk / Arbeitspaket","Beschreibung","Typ (F/E)","Anmerkung"],
    "1F3864", [12,8,34,52,12,40])

# ── 2_Projektplanung ──────────────────────────────────────────────────────
ws2 = wb.create_sheet("2_Projektplanung")
titelzeile(ws2,
    "PROJEKTPLANUNG – Logische Reihenfolge mit Abhängigkeiten",
    ["Nr.","Gewerk / Arbeitspaket","Abhängigkeit von","Typ",
     "Dauer (Wo)","Beginn (Wo)","Ende (Wo)","Begründung Fachfirma","Risiko"],
    "2E75B6", [6,36,28,8,10,12,10,52,32])

# ── 3_Meilensteine ────────────────────────────────────────────────────────
ws3 = wb.create_sheet("3_Meilensteine")
titelzeile(ws3,
    "MEILENSTEINPLAN – Kernsanierung EFH Baujahr 1969",
    ["Nr.","Meilenstein","Fertigstellungskriterien","Zeitpunkt\n(Monat)","Zuständig","Status","Notiz"],
    "C55A11", [6,36,52,14,20,16,30])

# ── 4_Gantt ───────────────────────────────────────────────────────────────
ws4 = wb.create_sheet("4_Gantt")
ws4.sheet_properties.tabColor = "375623"
ws4["A1"].value = "GANTT-DIAGRAMM – Kernsanierung EFH (24 Monate)"
ws4["A1"].font = Font(bold=True, color=WEISS, size=12)
ws4["A1"].fill = PatternFill("solid", fgColor=DUNKELBLAU)
ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws4.row_dimensions[1].height = 28
gantt_headers = ["Nr.","Gewerk / Arbeitspaket","Typ"] + [f"M{i}" for i in range(1,25)]
for i, h in enumerate(gantt_headers, 1):
    c = ws4.cell(row=2, column=i, value=h)
    c.font = Font(bold=True, color=WEISS, size=9)
    c.fill = PatternFill("solid", fgColor=DUNKELBLAU)
    c.alignment = Alignment(horizontal="center", vertical="center")
ws4.row_dimensions[2].height = 22
ws4.column_dimensions["A"].width = 6
ws4.column_dimensions["B"].width = 36
ws4.column_dimensions["C"].width = 8
for i in range(4, 28):
    ws4.column_dimensions[get_column_letter(i)].width = 4
ws4.merge_cells(f"A1:{get_column_letter(27)}1")
ws4.freeze_panes = "D3"

# ── 5_Budget ──────────────────────────────────────────────────────────────
ws5 = wb.create_sheet("5_Budget")
titelzeile(ws5,
    "BUDGETPLANUNG – Kernsanierung EFH Baujahr 1969 (Marktpreise 2024/2025)",
    ["Gewerk","Umfang","Eigen-\nleistung","Soll-Kosten\n(€)","Ist-Kosten\n(€)",
     "Differenz\n(€)","Förderung\nmöglich","KfW/BAFA\nProgramm","Kommentar"],
    "7030A0", [28,22,10,14,14,14,14,22,36])

# ── 6_Tracking ────────────────────────────────────────────────────────────
ws6 = wb.create_sheet("6_Tracking")
titelzeile(ws6,
    "PROJEKTTRACKING – Kostenüberwachung & Fortschritt",
    ["Gewerk","Firma / Person","Startdatum","Enddatum","Angebot (€)",
     "Auftrag (€)","Rechnung (€)","Differenz (€)","Fortschritt","Status","Erledigt","Notizen"],
    "C00000", [28,24,13,13,14,14,14,14,14,16,10,36])

# ── 7_Tipps ───────────────────────────────────────────────────────────────
ws7 = wb.create_sheet("7_Tipps")
titelzeile(ws7,
    "PRAXISTIPPS – Phase für Phase (für Bauherren ohne Fachkenntnisse)",
    ["Phase","Gewerk","Tipp / Empfehlung","Typischer Fehler","Einsparpotenzial","Risiko / Fallstrick"],
    "BF8F00", [14,22,52,42,30,40])

# ── 8_Kritischer_Pfad ─────────────────────────────────────────────────────
ws8 = wb.create_sheet("8_Kritischer_Pfad")
titelzeile(ws8,
    "KRITISCHER PFAD – Terminkritische Abhängigkeiten & Kaskadenrisiken",
    ["Rang","Gewerk","Vorgänger\n(Pflicht)","Puffer\n(Wo)","Kaskadeneffekt bei Verzug","Maßnahme bei Verzug","Priorität"],
    "833C11", [6,32,24,10,48,40,12])

# ── 9_Genehmigungen ───────────────────────────────────────────────────────
ws9 = wb.create_sheet("9_Genehmigungen")
titelzeile(ws9,
    "GENEHMIGUNGEN & VORSCHRIFTEN – Deutschland (GEG/VDE/TRGI/EEG/VOB)",
    ["Gewerk","Genehmigung erforderlich?","Zuständige Behörde","Norm / Vorschrift",
     "Konzessions-\npflicht","Abnahme durch","Frist / Hinweis","Förder-\nrelevanz"],
    "244185", [24,22,24,28,14,22,36,14])

# ── 10_Top10_Empfehlungen ─────────────────────────────────────────────────
ws10 = wb.create_sheet("10_Top10_Empfehlungen")
titelzeile(ws10,
    "TOP 10 EMPFEHLUNGEN – Das Wichtigste für Bauherren ohne Fachkenntnisse",
    ["Rang","Empfehlung","Begründung","Häufiger Fehler","Konsequenz bei Nichtbeachtung","Priorität"],
    "FF0000", [6,36,48,42,48,12])

wb.save("/home/user/Excelhaus/sanierungsplan.xlsx")
print("Grundgerüst erstellt: sanierungsplan.xlsx mit 10 Tabellenblättern")
